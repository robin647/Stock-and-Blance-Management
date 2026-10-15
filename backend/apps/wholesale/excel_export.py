from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import PermissionDenied
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

from .models import WholesaleInvoice


class WholesaleExcelExportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if request.user.role != "admin":
            raise PermissionDenied("Only Admin can export wholesale reports.")
        invoices = WholesaleInvoice.objects.select_related("party").prefetch_related("items__product", "items__returns", "payments")
        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")
        party_id = request.query_params.get("party")
        if start_date:
            invoices = invoices.filter(sale_date__gte=start_date)
        if end_date:
            invoices = invoices.filter(sale_date__lte=end_date)
        if party_id:
            invoices = invoices.filter(party_id=party_id)

        workbook = Workbook()
        invoice_sheet = workbook.active
        invoice_sheet.title = "Invoices"
        _header(invoice_sheet, ["Date", "Party", "Invoice", "Total Quantity", "Gross Sale", "Bank", "bKash",
                    "Cash", "Payments", "Returns", "Due", "Credit"])
        item_sheet = workbook.create_sheet("Invoice Items")
        _header(item_sheet, ["Invoice", "Product", "Code", "Size", "Qty", "Unit Price", "Line Total",
                             "Returned Qty", "Return Amount"])
        for invoice in invoices:
            payment_totals = [sum((getattr(payment, field) for payment in invoice.payments.all()), 0)
                              for field in ["bank_amount", "bkash_amount", "cash_amount"]]
            invoice_sheet.append([
                invoice.sale_date, invoice.party.name, invoice.invoice_number, invoice.total_quantity,
                invoice.gross_amount, *payment_totals, invoice.payment_total, invoice.return_total, invoice.due_amount,
                invoice.credit_amount,
            ])
            for item in invoice.items.all():
                item_sheet.append([
                    invoice.invoice_number, item.product.name, item.product.code, item.product.size, item.quantity,
                    item.unit_price, item.line_total,
                    sum(entry.quantity for entry in item.returns.all()),
                    sum((entry.amount for entry in item.returns.all()), 0),
                ])

        products = workbook.create_sheet("Product Quantities")
        _header(products, ["Product", "Code", "Quantity Sold", "Quantity Returned", "Return Amount"])
        product_rows = {}
        for invoice in invoices:
            for item in invoice.items.all():
                row = product_rows.setdefault(item.product_id, [item.product.name, item.product.code, 0, 0, 0])
                row[2] += item.quantity
                row[3] += sum(entry.quantity for entry in item.returns.all())
                row[4] += sum((entry.amount for entry in item.returns.all()), 0)
        for row in product_rows.values():
            products.append(row)

        party_sheet = workbook.create_sheet("Party Summary")
        _header(party_sheet, ["Party", "Gross Sales", "Payments", "Returns", "Outstanding Due", "Credit"])
        party_totals = {}
        for invoice in invoices:
            row = party_totals.setdefault(invoice.party_id, [invoice.party.name, 0, 0, 0, 0, 0])
            row[1] += invoice.gross_amount
            row[2] += invoice.payment_total
            row[3] += invoice.return_total
            row[4] += invoice.due_amount
            row[5] += invoice.credit_amount
        for row in party_totals.values():
            party_sheet.append(row)

        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="wholesale-report.xlsx"'
        workbook.save(response)
        return response


def _header(sheet, labels):
    sheet.append(labels)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315B52")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions