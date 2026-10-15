from decimal import Decimal
from datetime import date
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.inventory.models import FactoryStock, ShowroomStock, StockAdjustment
from apps.products.models import Product, ProductCategory
from apps.showrooms.models import Showroom
from .models import EasyWholesaleStockMovement


class WholesaleApiTests(TestCase):
    def setUp(self):
        category = ProductCategory.objects.create(name="Panjabi")
        self.product = Product.objects.create(
            name="Classic Panjabi", code="3333", category=category, size="42", selling_price=Decimal("120.00")
        )
        self.factory_stock = FactoryStock.objects.create(product=self.product, quantity=20)
        self.showroom = Showroom.objects.create(name="Zone 1")
        self.showroom_stock = ShowroomStock.objects.create(showroom=self.showroom, product=self.product, quantity=8)
        self.admin = User.objects.create_user(username="wholesale-admin", password="password", role="admin")
        self.showroom_user = User.objects.create_user(
            username="wholesale-showroom", password="password", role="showroom_user", showroom=self.showroom
        )
        self.client = APIClient()
        self.client.force_authenticate(self.admin)
        party_response = self.client.post("/api/wholesale/parties/", {"name": "Test Buyer"}, format="json")
        self.party_id = party_response.data["id"]

    def test_stock_adjustment_updates_only_selected_location_and_keeps_history(self):
        response = self.client.post("/api/wholesale/stock-adjustments/", {
            "product": self.product.id, "location": "factory", "new_quantity": 25, "reason": "Count correction",
        }, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.factory_stock.refresh_from_db()
        self.showroom_stock.refresh_from_db()
        adjustment = StockAdjustment.objects.get()
        self.assertEqual(self.factory_stock.quantity, 25)
        self.assertEqual(self.showroom_stock.quantity, 8)
        self.assertEqual(adjustment.previous_quantity, 20)
        self.assertEqual(adjustment.difference, 5)
        self.assertEqual(adjustment.reason, "Count correction")
        self.assertEqual(adjustment.created_by, self.admin)

        showroom_response = self.client.post("/api/wholesale/stock-adjustments/", {
            "product": self.product.id, "location": "showroom", "showroom": self.showroom.id,
            "new_quantity": 10, "reason": "Showroom count correction",
        }, format="json")
        self.assertEqual(showroom_response.status_code, 201, showroom_response.data)
        self.factory_stock.refresh_from_db()
        self.showroom_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 25)
        self.assertEqual(self.showroom_stock.quantity, 10)

    def test_invoice_payment_and_return_keep_quantities_and_due_consistent(self):
        invoice_response = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id,
            "invoice_number": "WS-001",
            "sale_date": "2026-10-15",
            "stock_location": "factory",
            "items": [{"product": self.product.id, "quantity": 3, "unit_price": "100.00"}],
            "initial_cash_payment": "50.00",
        }, format="json")

        self.assertEqual(invoice_response.status_code, 201, invoice_response.data)
        invoice_id = invoice_response.data["id"]
        invoice_item_id = invoice_response.data["items"][0]["id"]
        self.assertEqual(Decimal(invoice_response.data["gross_amount"]), Decimal("300.00"))
        self.assertEqual(invoice_response.data["total_quantity"], 3)
        self.assertEqual(Decimal(invoice_response.data["due_amount"]), Decimal("250.00"))
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 17)

        payment_response = self.client.post("/api/wholesale/payments/", {
            "invoice": invoice_id, "payment_date": "2026-10-15", "bank_amount": "100.00",
            "bkash_amount": "0.00", "cash_amount": "0.00",
        }, format="json")
        self.assertEqual(payment_response.status_code, 201, payment_response.data)

        return_response = self.client.post("/api/wholesale/returns/", {
            "invoice_item": invoice_item_id, "quantity": 1, "amount": "100.00", "return_date": "2026-10-15",
        }, format="json")
        self.assertEqual(return_response.status_code, 201, return_response.data)
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 18)

        detail = self.client.get(f"/api/wholesale/invoices/{invoice_id}/")
        self.assertEqual(Decimal(detail.data["total_payments"]), Decimal("150.00"))
        self.assertEqual(detail.data["return_quantity"], 1)
        self.assertEqual(Decimal(detail.data["return_amount"]), Decimal("100.00"))
        self.assertEqual(Decimal(detail.data["due_amount"]), Decimal("50.00"))

    def test_duplicate_invoice_and_excess_return_are_rejected(self):
        payload = {
            "party": self.party_id, "invoice_number": "WS-DUP", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }
        first = self.client.post("/api/wholesale/invoices/", payload, format="json")
        second = self.client.post("/api/wholesale/invoices/", payload, format="json")
        self.assertEqual(first.status_code, 201, first.data)
        self.assertEqual(second.status_code, 400)
        item_id = first.data["items"][0]["id"]
        response = self.client.post("/api/wholesale/returns/", {
            "invoice_item": item_id, "quantity": 3, "amount": "100.00", "return_date": "2026-10-15",
        }, format="json")
        self.assertEqual(response.status_code, 400)

    def test_showroom_invoice_deducts_and_return_restores_only_that_showroom(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-SHOWROOM", "sale_date": "2026-10-15",
            "stock_location": "showroom", "showroom": self.showroom.id,
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(invoice.status_code, 201, invoice.data)
        self.factory_stock.refresh_from_db()
        self.showroom_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)
        self.assertEqual(self.showroom_stock.quantity, 6)

        returned = self.client.post("/api/wholesale/returns/", {
            "invoice_item": invoice.data["items"][0]["id"], "quantity": 1,
            "amount": "100.00", "return_date": "2026-10-15",
        }, format="json")
        self.assertEqual(returned.status_code, 201, returned.data)
        self.factory_stock.refresh_from_db()
        self.showroom_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)
        self.assertEqual(self.showroom_stock.quantity, 7)

    def test_return_after_full_payment_is_tracked_as_credit_not_negative_due(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-CREDIT", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
            "initial_cash_payment": "200.00",
        }, format="json")
        self.assertEqual(invoice.status_code, 201, invoice.data)
        returned = self.client.post("/api/wholesale/returns/", {
            "invoice_item": invoice.data["items"][0]["id"], "quantity": 1,
            "amount": "100.00", "return_date": "2026-10-15",
        }, format="json")
        self.assertEqual(returned.status_code, 201, returned.data)
        detail = self.client.get(f"/api/wholesale/invoices/{invoice.data['id']}/")
        self.assertEqual(Decimal(detail.data["due_amount"]), Decimal("0.00"))
        self.assertEqual(Decimal(detail.data["credit_amount"]), Decimal("100.00"))

    def test_non_admin_cannot_manage_wholesale_or_stock_adjustments(self):
        self.client.force_authenticate(self.showroom_user)
        party = self.client.post("/api/wholesale/parties/", {"name": "Blocked Buyer"}, format="json")
        adjustment = self.client.post("/api/wholesale/stock-adjustments/", {
            "product": self.product.id, "location": "showroom", "showroom": self.showroom.id,
            "new_quantity": 12, "reason": "Unauthorized",
        }, format="json")
        summary = self.client.get("/api/wholesale/summary/")
        self.assertEqual(party.status_code, 403)
        self.assertEqual(adjustment.status_code, 403)
        self.assertEqual(summary.status_code, 403)

    def test_party_partial_search_edit_and_deactivation(self):
        search = self.client.get("/api/wholesale/parties/", {"search": "test buy"})
        self.assertEqual(search.status_code, 200)
        self.assertEqual(search.data["count"], 1)
        response = self.client.patch(f"/api/wholesale/parties/{self.party_id}/", {
            "contact_person": "A. Buyer", "is_active": False,
        }, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["contact_person"], "A. Buyer")
        self.assertFalse(response.data["is_active"])

    def test_party_list_endpoint_aggregates_legacy_and_easy_transactions(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "LIST-LEGACY", "sale_date": "2026-10-01",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(invoice.status_code, 201, invoice.data)
        payment = self.client.post("/api/wholesale/payments/", {
            "invoice": invoice.data["id"], "payment_date": "2026-10-02", "bank_amount": "50.00",
            "bkash_amount": "0.00", "cash_amount": "0.00",
        }, format="json")
        self.assertEqual(payment.status_code, 201, payment.data)
        returned = self.client.post("/api/wholesale/returns/", {
            "invoice_item": invoice.data["items"][0]["id"], "quantity": 1,
            "amount": "100.00", "return_date": "2026-10-03",
        }, format="json")
        self.assertEqual(returned.status_code, 201, returned.data)
        easy_sale = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-04", "bill_number": "LIST-EASY",
            "quantity": 1, "total_amount": "75.00",
        }, format="json")
        self.assertEqual(easy_sale.status_code, 201, easy_sale.data)

        response = self.client.get("/api/wholesale/parties/", {
            "search": "test", "is_active": "true", "page": 1, "page_size": 10,
        })
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["count"], 1)
        party_row = response.data["results"][0]
        self.assertEqual(Decimal(party_row["total_sales"]), Decimal("175.00"))
        self.assertEqual(Decimal(party_row["total_payments"]), Decimal("50.00"))
        self.assertEqual(Decimal(party_row["total_returns"]), Decimal("100.00"))
        self.assertEqual(Decimal(party_row["total_due"]), Decimal("125.00"))

    def test_party_routes_are_registered_with_trailing_slashes(self):
        self.assertEqual(reverse("wholesale-party-list"), "/api/wholesale/parties/")
        self.assertEqual(reverse("wholesale-party-detail", args=[1]), "/api/wholesale/parties/1/")

    def test_statement_uses_actual_payment_date_and_prior_opening_balance(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-DATE-01", "sale_date": "2026-10-01",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(invoice.status_code, 201, invoice.data)
        payment = self.client.post("/api/wholesale/payments/", {
            "invoice": invoice.data["id"], "payment_date": "2026-10-12", "bank_amount": "50.00",
            "bkash_amount": "25.00", "cash_amount": "25.00", "reference": "BANK-REF-7",
        }, format="json")
        self.assertEqual(payment.status_code, 201, payment.data)

        statement = self.client.get(f"/api/wholesale/parties/{self.party_id}/history/", {
            "start_date": "2026-10-10", "end_date": "2026-10-15",
        })
        self.assertEqual(statement.status_code, 200, statement.data)
        self.assertEqual(statement.data["opening_balance"], Decimal("200.00"))
        self.assertEqual(statement.data["summary"]["gross_sales"], Decimal("0.00"))
        self.assertEqual(statement.data["summary"]["total_payments"], Decimal("100.00"))
        self.assertEqual(statement.data["summary"]["bank_payments"], Decimal("50.00"))
        self.assertEqual(statement.data["closing_balance"], Decimal("100.00"))
        row = statement.data["results"][0]
        self.assertEqual(row["type"], "payment")
        self.assertEqual(row["date"], date(2026, 10, 12))
        self.assertEqual(row["reference"], "BANK-REF-7")
        self.assertEqual(row["running_balance"], Decimal("100.00"))

        statement_export = self.client.get(
            f"/api/wholesale/parties/{self.party_id}/statement/export/",
            {"start_date": "2026-10-10", "end_date": "2026-10-15"},
        )
        self.assertEqual(statement_export.status_code, 200)
        statement_sheet = load_workbook(BytesIO(statement_export.content), read_only=True)["Account Statement"]
        self.assertIn("BANK-REF-7", [cell for row in statement_sheet.iter_rows(values_only=True) for cell in row])

    def test_all_parties_export_uses_net_sales_and_preserves_summary(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-EXPORT", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(invoice.status_code, 201, invoice.data)
        self.client.post("/api/wholesale/returns/", {
            "invoice_item": invoice.data["items"][0]["id"], "quantity": 1,
            "amount": "100.00", "return_date": "2026-10-15",
        }, format="json")
        summary_export = self.client.get("/api/wholesale/summary/export/", {"search": "Test Buyer"})
        self.assertEqual(summary_export.status_code, 200)
        summary_sheet = load_workbook(BytesIO(summary_export.content), read_only=True)["All Parties Summary"]
        self.assertEqual(list(summary_sheet.iter_rows(values_only=True))[1][1], Decimal("100.00"))

    def test_failed_sale_rolls_back_invoice_and_stock(self):
        response = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-ROLLBACK", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 21, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.product.wholesale_sale_items.filter(invoice__invoice_number="WS-ROLLBACK").exists())
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)

    def test_payment_and_return_retries_are_idempotent(self):
        invoice = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "WS-IDEMPOTENT", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 2, "unit_price": "100.00"}],
        }, format="json")
        payment_payload = {
            "invoice": invoice.data["id"], "payment_date": "2026-10-15", "bank_amount": "25.00",
            "bkash_amount": "0.00", "cash_amount": "0.00",
        }
        first_payment = self.client.post("/api/wholesale/payments/", payment_payload, format="json", HTTP_IDEMPOTENCY_KEY="pay-retry-1")
        retry_payment = self.client.post("/api/wholesale/payments/", payment_payload, format="json", HTTP_IDEMPOTENCY_KEY="pay-retry-1")
        self.assertEqual(first_payment.data["id"], retry_payment.data["id"])

        item_id = invoice.data["items"][0]["id"]
        return_payload = {"invoice_item": item_id, "quantity": 1, "amount": "100.00", "return_date": "2026-10-15"}
        first_return = self.client.post("/api/wholesale/returns/", return_payload, format="json", HTTP_IDEMPOTENCY_KEY="return-retry-1")
        retry_return = self.client.post("/api/wholesale/returns/", return_payload, format="json", HTTP_IDEMPOTENCY_KEY="return-retry-1")
        self.assertEqual(first_return.data["id"], retry_return.data["id"])
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 19)
        detail = self.client.get(f"/api/wholesale/invoices/{invoice.data['id']}/")
        self.assertEqual(Decimal(detail.data["total_payments"]), Decimal("25.00"))
        self.assertEqual(Decimal(detail.data["return_amount"]), Decimal("100.00"))

    def test_easy_sale_without_product_saves_without_stock_movement_and_appears_in_statement(self):
        response = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-10", "bill_number": "EASY-100",
            "product_name": "Custom Sherwani", "quantity": 4, "total_amount": "1250.50",
            "notes": "Cash memo sale",
        }, format="json", HTTP_IDEMPOTENCY_KEY="easy-sale-test-1")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data["product"])
        self.factory_stock.refresh_from_db()
        self.showroom_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)
        self.assertEqual(self.showroom_stock.quantity, 8)
        self.assertEqual(EasyWholesaleStockMovement.objects.count(), 0)

        statement = self.client.get(f"/api/wholesale/parties/{self.party_id}/history/")
        self.assertEqual(statement.status_code, 200, statement.data)
        easy_rows = [row for row in statement.data["results"] if row.get("source") == "easy"]
        self.assertEqual(len(easy_rows), 1)
        self.assertEqual(easy_rows[0]["type"], "sale")
        self.assertEqual(easy_rows[0]["date"], date(2026, 10, 10))
        self.assertEqual(easy_rows[0]["product_name"], "Custom Sherwani")
        self.assertEqual(statement.data["summary"]["gross_sales"], Decimal("1250.50"))

    def test_easy_party_payment_is_a_separate_dated_statement_event(self):
        sale = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-01", "bill_number": "EASY-101",
            "quantity": 2, "total_amount": "500.00",
        }, format="json")
        self.assertEqual(sale.status_code, 201, sale.data)
        payment = self.client.post("/api/wholesale/easy/payments/", {
            "party": self.party_id, "payment_date": "2026-10-12", "bank_amount": "20.00",
            "bkash_amount": "30.00", "cash_amount": "50.00", "reference": "EASY-REF",
        }, format="json", HTTP_IDEMPOTENCY_KEY="easy-payment-test-1")
        self.assertEqual(payment.status_code, 201, payment.data)
        self.assertEqual(payment.data["total_amount"], "100.00")

        filtered = self.client.get(f"/api/wholesale/parties/{self.party_id}/history/", {
            "start_date": "2026-10-10", "end_date": "2026-10-15", "transaction_type": "payment",
        })
        self.assertEqual(filtered.status_code, 200, filtered.data)
        self.assertEqual(filtered.data["opening_balance"], Decimal("500.00"))
        self.assertEqual(filtered.data["summary"]["gross_sales"], Decimal("0.00"))
        self.assertEqual(filtered.data["summary"]["bank_payments"], Decimal("20.00"))
        self.assertEqual(filtered.data["summary"]["bkash_payments"], Decimal("30.00"))
        self.assertEqual(filtered.data["summary"]["cash_payments"], Decimal("50.00"))
        self.assertEqual(filtered.data["closing_balance"], Decimal("400.00"))
        self.assertEqual(filtered.data["results"][0]["date"], date(2026, 10, 12))
        self.assertEqual(filtered.data["results"][0]["reference"], "EASY-REF")

    def test_easy_standalone_return_is_financial_only_and_linked_return_restores_selected_stock(self):
        standalone = self.client.post("/api/wholesale/easy/returns/", {
            "party": self.party_id, "return_date": "2026-10-11", "quantity": 1,
            "amount": "40.00", "product_name": "Old Sherwani", "notes": "No physical item returned",
        }, format="json")
        self.assertEqual(standalone.status_code, 201, standalone.data)
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)

        sale = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-09", "bill_number": "EASY-102",
            "product": self.product.id, "quantity": 3, "total_amount": "360.00",
            "stock_location": "showroom", "showroom": self.showroom.id,
        }, format="json")
        self.assertEqual(sale.status_code, 201, sale.data)
        payment = self.client.post("/api/wholesale/easy/payments/", {
            "party": self.party_id, "payment_date": "2026-10-10", "bank_amount": "10.00",
            "bkash_amount": "0.00", "cash_amount": "0.00",
        }, format="json")
        self.assertEqual(payment.status_code, 201, payment.data)
        returned = self.client.post("/api/wholesale/easy/returns/", {
            "party": self.party_id, "return_date": "2026-10-13", "easy_sale": sale.data["id"],
            "quantity": 1, "amount": "120.00",
        }, format="json")
        self.assertEqual(returned.status_code, 201, returned.data)
        self.assertEqual(EasyWholesaleStockMovement.objects.count(), 2)
        self.showroom_stock.refresh_from_db()
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.showroom_stock.quantity, 6)
        self.assertEqual(self.factory_stock.quantity, 20)
        movements = list(EasyWholesaleStockMovement.objects.order_by("movement_type"))
        self.assertEqual(len(movements), 2)
        self.assertEqual({movement.movement_type for movement in movements}, {"sale", "return"})
        self.assertTrue(all(movement.stock_location == "showroom" for movement in movements))

        statement = self.client.get(f"/api/wholesale/parties/{self.party_id}/history/")
        self.assertEqual(statement.status_code, 200, statement.data)
        types = [row["type"] for row in statement.data["results"]]
        self.assertIn("sale", types)
        self.assertIn("payment", types)
        self.assertIn("return", types)

    def test_easy_routes_are_admin_only_and_payment_requires_positive_total(self):
        invalid = self.client.post("/api/wholesale/easy/payments/", {
            "party": self.party_id, "payment_date": "2026-10-15", "bank_amount": "0",
            "bkash_amount": "0", "cash_amount": "0",
        }, format="json")
        self.assertEqual(invalid.status_code, 400)
        self.client.force_authenticate(self.showroom_user)
        response = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-15", "quantity": 1, "total_amount": "1.00",
        }, format="json")
        self.assertEqual(response.status_code, 403)

    def test_easy_transactions_are_included_once_in_party_summary_and_legacy_sales_remain_available(self):
        legacy = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "LEGACY-200", "sale_date": "2026-10-10",
            "items": [{"product": self.product.id, "quantity": 1, "unit_price": "100.00"}],
        }, format="json")
        self.assertEqual(legacy.status_code, 201, legacy.data)
        easy = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-11", "bill_number": "EASY-200",
            "quantity": 2, "total_amount": "250.00",
        }, format="json")
        self.assertEqual(easy.status_code, 201, easy.data)

        legacy_records = self.client.get("/api/wholesale/invoices/", {"party": self.party_id})
        self.assertEqual(legacy_records.status_code, 200)
        self.assertEqual(legacy_records.data["count"], 1)
        summary = self.client.get("/api/wholesale/summary/", {"party": self.party_id})
        self.assertEqual(summary.status_code, 200, summary.data)
        party_row = summary.data["results"][0]
        self.assertEqual(Decimal(party_row["total_sales"]), Decimal("350.00"))
        self.assertEqual(party_row["total_products"], 3)

    def test_easy_product_sale_deducts_only_when_explicit_inventory_location_is_selected(self):
        financial = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-15", "quantity": 1,
            "product": self.product.id, "product_name": self.product.name, "total_amount": "120.00",
        }, format="json")
        self.assertEqual(financial.status_code, 201, financial.data)
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 20)
        self.assertEqual(EasyWholesaleStockMovement.objects.count(), 0)

        physical = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-15", "bill_number": "EASY-STOCK-01",
            "quantity": 2, "product": self.product.id, "total_amount": "240.00",
            "stock_location": "factory",
        }, format="json")
        self.assertEqual(physical.status_code, 201, physical.data)
        self.factory_stock.refresh_from_db()
        self.assertEqual(self.factory_stock.quantity, 18)
        movement = EasyWholesaleStockMovement.objects.get(easy_sale_id=physical.data["id"])
        self.assertEqual(movement.movement_type, "sale")
        self.assertEqual(movement.quantity, 2)

    def test_legacy_and_easy_statement_rows_with_same_table_ids_are_not_collapsed(self):
        legacy = self.client.post("/api/wholesale/invoices/", {
            "party": self.party_id, "invoice_number": "LEGACY-SAME-ID", "sale_date": "2026-10-15",
            "items": [{"product": self.product.id, "quantity": 1, "unit_price": "100.00"}],
        }, format="json")
        easy = self.client.post("/api/wholesale/easy/sales/", {
            "party": self.party_id, "sale_date": "2026-10-15", "bill_number": "EASY-SAME-ID",
            "quantity": 1, "total_amount": "250.00",
        }, format="json")
        self.assertEqual(legacy.status_code, 201, legacy.data)
        self.assertEqual(easy.status_code, 201, easy.data)
        self.assertEqual(legacy.data["id"], easy.data["id"])

        statement = self.client.get(f"/api/wholesale/parties/{self.party_id}/history/")
        self.assertEqual(statement.status_code, 200, statement.data)
        rows = [row for row in statement.data["results"] if row["type"] == "sale"]
        self.assertEqual(len(rows), 2)
        self.assertEqual({row["source"] for row in rows}, {"legacy", "easy"})
        self.assertEqual([row["running_balance"] for row in rows], [Decimal("100.00"), Decimal("350.00")])