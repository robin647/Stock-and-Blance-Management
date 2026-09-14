from datetime import date

from django.db.models import Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from rest_framework import permissions
from rest_framework.views import APIView

from apps.inventory.models import DailyBalance, DailyStock

from .helpers import resolve_showroom_scope


def _styled_header(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


class DailyExcelExportView(APIView):
    """
    GET /api/reports/export/daily/?date=YYYY-MM-DD&showroom=<id>
    Downloads one showroom's daily statement (stock + balance) as .xlsx.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        target_date = request.query_params.get("date", date.today().isoformat())
        showrooms_qs, showroom_id = resolve_showroom_scope(request)
        if not showroom_id:
            return HttpResponse("showroom parameter is required", status=400)

        wb = Workbook()

        ws_stock = wb.active
        ws_stock.title = "Stock"
        _styled_header(ws_stock, ["Product", "Opening", "Received", "Sold", "Return", "Closing"])
        stocks = DailyStock.objects.filter(showroom_id=showroom_id, date=target_date).select_related("product")
        for s in stocks:
            ws_stock.append([s.product.name, s.opening_qty, s.received_qty, s.sold_qty, s.return_qty, s.closing_qty])

        ws_balance = wb.create_sheet("Balance")
        _styled_header(ws_balance, ["Item", "Amount"])
        balance = DailyBalance.objects.filter(showroom_id=showroom_id, date=target_date).first()
        if balance:
            rows = [
                ("Opening Balance", balance.opening_balance),
                ("Cash Sale", balance.cash_sale),
                ("Card Sale", balance.card_sale),
                ("Total Sale", balance.total_sale),
                ("Expense", balance.expense),
                ("Salary", balance.salary),
                ("Deposit", balance.deposit),
                ("Closing Balance", balance.closing_balance),
            ]
            for row in rows:
                ws_balance.append(row)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"daily-statement-{showroom_id}-{target_date}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


class MonthlyExcelExportView(APIView):
    """
    GET /api/reports/export/monthly/?year=YYYY&month=MM&showroom=<id optional for admin>
    Downloads the monthly balance + stock summary as .xlsx.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        year = int(request.query_params.get("year", date.today().year))
        month = int(request.query_params.get("month", date.today().month))
        showrooms_qs, showroom_id = resolve_showroom_scope(request)

        balances = DailyBalance.objects.filter(date__year=year, date__month=month, showroom__in=showrooms_qs)
        if showroom_id:
            balances = balances.filter(showroom_id=showroom_id)

        stocks = DailyStock.objects.filter(date__year=year, date__month=month, showroom__in=showrooms_qs)
        if showroom_id:
            stocks = stocks.filter(showroom_id=showroom_id)

        wb = Workbook()
        ws_balance = wb.active
        ws_balance.title = "Balance Summary"
        _styled_header(ws_balance, ["Showroom", "Total Sale", "Cash Sale", "Card Sale", "Expense", "Salary", "Deposit"])
        balance_summary = balances.values("showroom__name").annotate(
            total_sale=Sum("total_sale"), cash_sale=Sum("cash_sale"), card_sale=Sum("card_sale"),
            expense=Sum("expense"), salary=Sum("salary"), deposit=Sum("deposit"),
        )
        for row in balance_summary:
            ws_balance.append([
                row["showroom__name"], row["total_sale"], row["cash_sale"], row["card_sale"],
                row["expense"], row["salary"], row["deposit"],
            ])

        ws_stock = wb.create_sheet("Stock Summary")
        _styled_header(ws_stock, ["Product", "Received", "Sold", "Return"])
        stock_summary = stocks.values("product__name").annotate(
            received=Sum("received_qty"), sold=Sum("sold_qty"), returned=Sum("return_qty"),
        )
        for row in stock_summary:
            ws_stock.append([row["product__name"], row["received"], row["sold"], row["returned"]])

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = f"monthly-report-{year}-{month:02d}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response
